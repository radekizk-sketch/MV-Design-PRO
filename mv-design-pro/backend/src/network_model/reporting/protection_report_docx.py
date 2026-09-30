"""
FIX-12: Protection Coordination DOCX Report Generator

Creates Word document report with:
- Summary numbers (smallest grading margin and ratios — no verdict, P-06)
- Device settings table
- Sensitivity/selectivity/overload check tables
- TCC data reference
- Binary determinism (same input → identical bytes)

CANONICAL ALIGNMENT:
- NOT-A-SOLVER: Only formatting, no physics calculations
- 100% Polish labels
- Deterministic output (uses shared docx_determinism module)
"""

from __future__ import annotations

import math
from io import BytesIO
from pathlib import Path
from typing import Any

from network_model.nazwy import jest_nazwa

# Import shared determinism module
from network_model.reporting.docx_determinism import make_docx_bytes_deterministic
from network_model.reporting.protection_tcc_presentation import (
    NAGLOWKI_NASTAW_PL,
    etykieta_tms,
    etykieta_typu_krzywej_pl,
    nazwa_urzadzenia,
    nazwa_wpisu_urzadzenia,
    nazwy_urzadzen,
    powod_braku_pl,
    uzasadnienia_sprawdzen,
    wiersze_nastaw_urzadzenia,
    wiersze_podsumowania_pl,
)

# Check for python-docx availability
try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False


def _format_value(value: Any) -> str:
    """Format a value for display."""
    if value is None:
        return "—"
    if isinstance(value, float):
        # Wynik nie niesie liczb zastępczych (brak = ``None``); ∞ wyłącznie dla nieskończoności.
        if math.isinf(value):
            return "∞"
        return f"{value:.3f}"
    if isinstance(value, bool):
        return "Tak" if value else "Nie"
    return str(value)


def export_protection_coordination_to_docx(
    result: dict[str, Any],
    path: str | Path,
    *,
    title: str | None = None,
    metadata: dict[str, Any] | None = None,
    deterministic: bool = True,
) -> Path:
    """
    Export protection coordination result to DOCX.

    Args:
        result: Coordination analysis result dict
        path: Target file path
        title: Custom report title
        metadata: Optional metadata (project_name, created_at, etc.)
        deterministic: If True, ensure binary reproducibility

    Returns:
        Path to the written DOCX file

    Raises:
        ImportError: If python-docx is not installed
    """
    if not _DOCX_AVAILABLE:
        raise ImportError("DOCX export requires python-docx. Install with: pip install python-docx")

    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # Create document
    doc = Document()

    # ==========================================================================
    # TITLE
    # ==========================================================================
    report_title = title if title else "Raport koordynacji zabezpieczeń nadprądowych"
    title_para = doc.add_heading(report_title, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Metadata — koordynacja nie wydaje werdyktu ogólnego (P-06), tylko liczby niżej.
    if metadata:
        meta_para = doc.add_paragraph()
        if jest_nazwa(metadata.get("project_name")):
            meta_para.add_run(f"Projekt: {metadata['project_name']}  |  ")
        if metadata.get("created_at"):
            meta_para.add_run(f"Data: {metadata['created_at'][:19]}")
        meta_para.paragraph_format.space_after = Pt(12)

    doc.add_paragraph()

    # ==========================================================================
    # SUMMARY
    # ==========================================================================
    doc.add_heading("Podsumowanie", level=1)

    summary_table = doc.add_table(rows=1, cols=2)
    summary_table.style = "Table Grid"

    # Header
    hdr_cells = summary_table.rows[0].cells
    hdr_cells[0].text = "Parametr"
    hdr_cells[1].text = "Wartość"
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].bold = True

    # Data rows
    for label, value in wiersze_podsumowania_pl(result["summary"], _format_value):
        row = summary_table.add_row().cells
        row[0].text = label.rstrip(":")
        row[1].text = value

    doc.add_paragraph()

    # ==========================================================================
    # DEVICE SETTINGS TABLE
    # ==========================================================================
    doc.add_heading("Tabela urządzeń i nastaw", level=1)

    devices = result.get("devices", [])
    # Tabele sprawdzeń nazywają urządzenie nazwą z listy urządzeń wyniku (karta #144).
    nazwy = nazwy_urzadzen(result)
    if devices:
        # Sort by device name for deterministic order
        sorted_devices = sorted(
            devices, key=lambda d: (nazwa_wpisu_urzadzenia(d), str(d.get("id", "")))
        )

        dev_table = doc.add_table(rows=1, cols=len(NAGLOWKI_NASTAW_PL))
        dev_table.style = "Table Grid"

        hdr_cells = dev_table.rows[0].cells
        for i, h in enumerate(NAGLOWKI_NASTAW_PL):
            hdr_cells[i].text = h
            hdr_cells[i].paragraphs[0].runs[0].bold = True

        for dev in sorted_devices:
            for wiersz in wiersze_nastaw_urzadzenia(dev):
                row = dev_table.add_row().cells
                row[0].text = wiersz[0]
                for i, wartosc in enumerate(wiersz[1:], start=1):
                    row[i].text = wartosc
    else:
        doc.add_paragraph("Brak urządzeń", style="No Spacing")

    doc.add_paragraph()

    # ==========================================================================
    # SENSITIVITY CHECKS
    # ==========================================================================
    doc.add_heading("Czułość — iloraz I_min / I_s", level=1)

    sensitivity_checks = result.get("sensitivity_checks", [])
    if sensitivity_checks:
        # Sort by device_id for deterministic order
        sorted_sens_checks = sorted(sensitivity_checks, key=lambda c: c.get("device_id", ""))

        sens_table = doc.add_table(rows=1, cols=5)
        sens_table.style = "Table Grid"

        # Header
        headers = ["Urządzenie", "I_min [A]", "I_s [A]", "Iloraz", "Wymagany"]
        hdr_cells = sens_table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].paragraphs[0].runs[0].bold = True

        # Data
        for check in sorted_sens_checks:
            row = sens_table.add_row().cells
            row[0].text = nazwa_urzadzenia(nazwy, check["device_id"])
            row[1].text = _format_value(check["i_fault_min_a"])
            row[2].text = _format_value(check["i_pickup_a"])
            row[3].text = _format_value(check["ratio"])
            row[4].text = _format_value(check["required_ratio"])
        for zdanie in uzasadnienia_sprawdzen(sorted_sens_checks, nazwy):
            doc.add_paragraph(zdanie, style="No Spacing")
    else:
        doc.add_paragraph("Brak danych", style="No Spacing")

    doc.add_paragraph()

    # ==========================================================================
    # SELECTIVITY CHECKS
    # ==========================================================================
    doc.add_heading("Selektywność czasowa — odstęp t_nad − t_pod", level=1)

    selectivity_checks = result.get("selectivity_checks", [])
    if selectivity_checks:
        # Sort by (downstream_device_id, upstream_device_id) for deterministic order
        sorted_sel_checks = sorted(
            selectivity_checks,
            key=lambda c: (c.get("downstream_device_id", ""), c.get("upstream_device_id", "")),
        )

        sel_table = doc.add_table(rows=1, cols=6)
        sel_table.style = "Table Grid"

        # Header
        headers = ["Podrzędne", "Nadrzędne", "t_pod [s]", "t_nad [s]", "Odstęp [s]", "Wymagany [s]"]
        hdr_cells = sel_table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].paragraphs[0].runs[0].bold = True

        # Data
        for check in sorted_sel_checks:
            row = sel_table.add_row().cells
            row[0].text = nazwa_urzadzenia(nazwy, check["downstream_device_id"])
            row[1].text = nazwa_urzadzenia(nazwy, check["upstream_device_id"])
            row[2].text = _format_value(check["t_downstream_s"])
            row[3].text = _format_value(check["t_upstream_s"])
            row[4].text = _format_value(check["margin_s"])
            row[5].text = _format_value(check["required_margin_s"])
        for zdanie in uzasadnienia_sprawdzen(sorted_sel_checks, nazwy):
            doc.add_paragraph(zdanie, style="No Spacing")
    else:
        doc.add_paragraph(
            "Brak par stopniowania (strefy urządzeń nie są zagnieżdżone albo para jest nierozstrzygalna)",
            style="No Spacing",
        )

    doc.add_paragraph()

    # ==========================================================================
    # OVERLOAD CHECKS
    # ==========================================================================
    doc.add_heading("Przeciążalność — iloraz I_s / I_rob", level=1)

    overload_checks = result.get("overload_checks", [])
    if overload_checks:
        # Sort by device_id for deterministic order
        sorted_ovl_checks = sorted(overload_checks, key=lambda c: c.get("device_id", ""))

        ovl_table = doc.add_table(rows=1, cols=5)
        ovl_table.style = "Table Grid"

        # Header
        headers = ["Urządzenie", "I_rob [A]", "I_s [A]", "Iloraz", "Wymagany"]
        hdr_cells = ovl_table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].paragraphs[0].runs[0].bold = True

        # Data
        for check in sorted_ovl_checks:
            row = ovl_table.add_row().cells
            row[0].text = nazwa_urzadzenia(nazwy, check["device_id"])
            row[1].text = _format_value(check["i_operating_a"])
            row[2].text = _format_value(check["i_pickup_a"])
            row[3].text = _format_value(check["ratio"])
            row[4].text = _format_value(check["required_ratio"])
        for zdanie in uzasadnienia_sprawdzen(sorted_ovl_checks, nazwy):
            doc.add_paragraph(zdanie, style="No Spacing")
    else:
        doc.add_paragraph("Brak danych", style="No Spacing")

    doc.add_paragraph()

    # ==========================================================================
    # TCC CURVES INFO
    # ==========================================================================
    doc.add_heading("Krzywe czasowo-prądowe (TCC)", level=1)

    tcc_curves = result.get("tcc_curves", [])
    if tcc_curves:
        # Sort by device_name for deterministic order
        sorted_tcc_curves = sorted(tcc_curves, key=lambda c: c.get("device_name", ""))

        _powody_bez_podstawy: list[str] = []
        tcc_table = doc.add_table(rows=1, cols=4)
        tcc_table.style = "Table Grid"

        # Header
        headers = ["Urządzenie", "Typ krzywej", "I_s [A]", "TMS"]
        hdr_cells = tcc_table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].paragraphs[0].runs[0].bold = True

        # Data
        for curve in sorted_tcc_curves:
            row = tcc_table.add_row().cells
            row[0].text = nazwa_wpisu_urzadzenia(curve, "device_name")
            row[1].text = etykieta_typu_krzywej_pl(curve)
            row[2].text = _format_value(curve.get("pickup_current_a"))
            row[3].text = etykieta_tms(curve, _format_value(curve.get("time_multiplier")))
            # Pozycja bez podstawy (np. bezpiecznik bez pasma topikowego) niesie
            # PELNE zdanie po polsku pod tabela — ta sama tresc co w PDF, z tego
            # samego modulu prezentacji (karta N-D5-FUSE).
            powod = powod_braku_pl(curve)
            if powod:
                _powody_bez_podstawy.append(
                    f"{nazwa_wpisu_urzadzenia(curve, 'device_name')}: {powod}"
                )

        for powod in _powody_bez_podstawy:
            doc.add_paragraph(powod, style="No Spacing")

        doc.add_paragraph()
        note = doc.add_paragraph()
        note.add_run("Wykres TCC dostępny w eksporcie interaktywnym").italic = True
    else:
        doc.add_paragraph("Brak krzywych TCC", style="No Spacing")

    doc.add_paragraph()

    # ==========================================================================
    # FOOTER
    # ==========================================================================
    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.add_run(f"Identyfikator obliczenia: {result['run_id']}").font.size = Pt(8)
    footer.add_run(f"  |  Wygenerowano: {result.get('created_at', '—')[:19]}").font.size = Pt(8)

    # Save to bytes
    buffer = BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()

    # Make deterministic if requested (use shared module)
    if deterministic:
        docx_bytes = make_docx_bytes_deterministic(docx_bytes)

    # Write to file
    output_path.write_bytes(docx_bytes)
    return output_path
