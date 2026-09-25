"""Sekcja zgodności NC RfG w dokumentach formalnych — JEDEN serializer dla certyfikatu zgodności,
wniosku do OSD i dokumentu studium przyłączeniowego (karta AB-1a Pakiet C).

Dokumenty nie mają własnej oceny ani własnego tekstu werdyktu: czytają rekordy
``WynikWymagania`` z oceny wymagań (``application.ncrfg_compliance.ocena_wymagan``) i serializują
je ``werdykt.blok_wymagania`` (pozycja → wiersz). Poza rekordami sekcja modułu cytuje wyłącznie
dane, na których rekordy stoją: klasyfikację modułu z powodem i podstawą, technologię oraz dowód
certyfikatu urządzenia wyprowadzony przez serwer z wykazu PTPiREE (albo powód odrzucenia
tabliczki). Podstawy cytuje ``werdykt.opis_podstawy`` — ten sam format co w blokach rekordów.

Wiersze sekcji (``wiersze_sekcji``) to JEDNA lista, z której czerpią renderery DOCX i PDF —
zdanie rekordu w JSON, w DOCX i w PDF jest to samo (kontrakt werdyktu, T13).

Warstwa APPLICATION: zero fizyki, zero statusu, zero mapy status → tekst.
"""

from __future__ import annotations

from typing import Any

from application.ncrfg_compliance import (
    NcRfgCaseComplianceResponse,
    NcRfgCertyfikatOdrzucony,
    NcRfgDerPominiety,
    OcenaWymaganModulu,
)
from network_model.nazwy import nazwa_nadana
from network_model.solvers.ncrfg_ptpiree.contracts import (
    DowodCertyfikatu,
    NcRfgPtpireeModuleResult,
)
from network_model.solvers.ncrfg_ptpiree.stosowalnosc import NAZWA_TECHNOLOGII_PL
from werdykt import PozycjaBloku, blok_wymagania, format_liczba, opis_podstawy
from werdykt.dokument import POZYCJA_KRYTERIUM_SKLADOWE, POZYCJA_WYMAGANIE

TYTUL_SEKCJI_MODULU = "Moduł wytwarzania energii"
POZYCJA_MODUL = "Moduł"
POZYCJA_KLASYFIKACJA = "Klasyfikacja modułu"
POZYCJA_PODSTAWA_KLASYFIKACJI = "Podstawa klasyfikacji"
POZYCJA_TECHNOLOGIA = "Technologia"
POZYCJA_DOWOD_CERTYFIKATU = "Dowód certyfikatu urządzenia"
POZYCJA_PODSTAWA_DOWODU = "Podstawa dowodu certyfikatu"
POZYCJA_CERTYFIKAT_ODRZUCONY = "Certyfikat z tabliczki odrzucony"
POZYCJA_POMINIETY = "Źródło nieobjęte oceną"

#: Uczciwy stan zerowy dowodu certyfikatu (tabliczka nie wskazuje rekordu wykazu).
BRAK_CERTYFIKATU_PL = (
    "tabliczka urządzenia w modelu nie wskazuje rekordu wykazu certyfikowanych urządzeń PTPiREE"
)


def opis_dowodu_certyfikatu(dowod: DowodCertyfikatu) -> str:
    """Cytat rekordu wykazu PTPiREE (pola z rejestru, nie z tabliczki)."""
    czesci = [
        f"rekord wykazu {dowod.rekord_id}",
        f"producent: {dowod.producent}",
        f"model: {dowod.model}",
        f"numer dokumentu: {dowod.numer_dokumentu}",
        f"data akceptacji: {dowod.data_akceptacji or 'nie podano w wykazie'}",
        f"wersja WiPWC: {dowod.wersja_wipwc}",
        f"wersja WOS: {dowod.wersja_wos or 'nie podano w wykazie'}",
        "zakres typów modułów: "
        + (", ".join(dowod.zakres_typow) if dowod.zakres_typow else "pusty"),
    ]
    if dowod.warunek_waznosci is not None:
        czesci.append(f"warunek ważności: {dowod.warunek_waznosci}")
    if dowod.adres_zrodla is not None:
        czesci.append(f"źródło: {dowod.adres_zrodla}")
    return "; ".join(czesci)


def wiersze_dowodu(
    dowod: DowodCertyfikatu | None, odrzucony: NcRfgCertyfikatOdrzucony | None
) -> list[PozycjaBloku]:
    """Wiersze dowodu certyfikatu urządzenia: rekord wykazu z podstawą, powód odrzucenia
    tabliczki albo jawny stan zerowy."""
    if dowod is not None:
        return [
            PozycjaBloku(
                etykieta_pl=POZYCJA_DOWOD_CERTYFIKATU, tresc_pl=opis_dowodu_certyfikatu(dowod)
            ),
            PozycjaBloku(
                etykieta_pl=POZYCJA_PODSTAWA_DOWODU, tresc_pl=opis_podstawy(dowod.podstawa)
            ),
        ]
    if odrzucony is not None:
        return [PozycjaBloku(etykieta_pl=POZYCJA_CERTYFIKAT_ODRZUCONY, tresc_pl=odrzucony.powod_pl)]
    return [PozycjaBloku(etykieta_pl=POZYCJA_DOWOD_CERTYFIKATU, tresc_pl=BRAK_CERTYFIKATU_PL)]


def _nazwa_modulu(der_name: str | None) -> str:
    """Nazwa źródła DER w dokumencie: nazwa z modelu albo opis rodzaju — identyfikator
    modułu (`der_ref`) nie jest nazwą pokazywaną projektantowi (karta #144)."""
    return nazwa_nadana(der_name) or "Źródło DER bez nazwy"


def _wiersze_modulu(
    modul: NcRfgPtpireeModuleResult, odrzucony: NcRfgCertyfikatOdrzucony | None
) -> list[PozycjaBloku]:
    return [
        PozycjaBloku(
            etykieta_pl=POZYCJA_MODUL,
            tresc_pl=(
                f"{_nazwa_modulu(modul.der_name)}; operator: "
                f"{modul.operator_name_pl}; moc maksymalna {format_liczba(modul.p_max_kw)} kW; "
                f"napięcie przyłączenia {format_liczba(modul.voltage_kv)} kV"
            ),
        ),
        PozycjaBloku(etykieta_pl=POZYCJA_KLASYFIKACJA, tresc_pl=modul.klasyfikacja.powod_pl),
        PozycjaBloku(
            etykieta_pl=POZYCJA_PODSTAWA_KLASYFIKACJI,
            tresc_pl=opis_podstawy(modul.klasyfikacja.podstawa),
        ),
        PozycjaBloku(
            etykieta_pl=POZYCJA_TECHNOLOGIA, tresc_pl=NAZWA_TECHNOLOGII_PL[modul.technologia]
        ),
        *wiersze_dowodu(modul.dowod_certyfikatu, odrzucony),
    ]


def sekcja_modulu(
    modul: NcRfgPtpireeModuleResult,
    ocena: OcenaWymaganModulu,
    odrzucony: NcRfgCertyfikatOdrzucony | None,
) -> dict[str, Any]:
    """Sekcja jednego modułu: dane, na których stoją rekordy, i rekordy W z blokami."""
    if modul.der_ref != ocena.der_ref:
        raise ValueError(f"Wynik modułu {modul.der_ref} i ocena {ocena.der_ref} rozjechane.")
    return {
        "der_ref": modul.der_ref,
        "der_name": _nazwa_modulu(modul.der_name),
        "operator_pl": modul.operator_name_pl,
        "p_max_kw": modul.p_max_kw,
        "voltage_kv": modul.voltage_kv,
        "klasyfikacja": modul.klasyfikacja.model_dump(mode="json"),
        "technologia": modul.technologia,
        "zrodlo_danych": modul.zrodlo_danych,
        "dowod_certyfikatu": (
            modul.dowod_certyfikatu.model_dump(mode="json")
            if modul.dowod_certyfikatu is not None
            else None
        ),
        "certyfikat_odrzucony": (
            odrzucony.model_dump(mode="json") if odrzucony is not None else None
        ),
        "wiersze": [p.model_dump(mode="json") for p in _wiersze_modulu(modul, odrzucony)],
        "wymagania": [
            {
                "rekord": rekord.model_dump(mode="json"),
                "blok": [p.model_dump(mode="json") for p in blok_wymagania(rekord)],
            }
            for rekord in ocena.wymagania
        ],
    }


def sekcje_modulow(zgodnosc: NcRfgCaseComplianceResponse) -> list[dict[str, Any]]:
    """Sekcje wszystkich modułów biegu zgodności przypadku (kolejność modułów modelu)."""
    if zgodnosc.bieg is None:
        return []
    odrzucone = {o.der_ref: o for o in zgodnosc.certyfikaty_odrzucone}
    return [
        sekcja_modulu(modul, ocena, odrzucone.get(modul.der_ref))
        for modul, ocena in zip(zgodnosc.bieg.modules, zgodnosc.bieg.ocena_wymagan, strict=True)
    ]


def opis_pominietego(pominiety: NcRfgDerPominiety) -> str:
    """Źródło modelu nieobjęte oceną — nazwa i powód (stan zerowy per urządzenie)."""
    return f"{_nazwa_modulu(pominiety.der_name)}: {pominiety.powod_pl}"


def wiersze_sekcji(sekcja: dict[str, Any]) -> list[tuple[str, str, int]]:
    """Płaska lista wierszy sekcji modułu dla rendererów (etykieta, treść, poziom).

    Poziom 0 — dane modułu; poziom 1 — pozycje bloku wymagania; poziom 2 — pozycje bloku
    kryterium składowego (po wierszu „Kryterium składowe"). Kolejność = kolejność bloków.
    """
    wiersze = [(w["etykieta_pl"], w["tresc_pl"], 0) for w in sekcja["wiersze"]]
    for wymaganie in sekcja["wymagania"]:
        poziom = 1
        for pozycja in wymaganie["blok"]:
            if pozycja["etykieta_pl"] == POZYCJA_KRYTERIUM_SKLADOWE:
                wiersze.append((pozycja["etykieta_pl"], pozycja["tresc_pl"], 1))
                poziom = 2
                continue
            if pozycja["etykieta_pl"] == POZYCJA_WYMAGANIE:
                poziom = 1
            wiersze.append((pozycja["etykieta_pl"], pozycja["tresc_pl"], poziom))
    return wiersze


def dopisz_sekcje_docx(doc: Any, sekcje: list[dict[str, Any]], *, poziom_naglowka: int) -> None:
    """Dopisz sekcje modułów do dokumentu DOCX (``python-docx``): nagłówek modułu, potem
    wiersze ``wiersze_sekcji`` jako akapity „etykieta: treść" z wcięciem wg poziomu."""
    from docx.shared import Pt

    for sekcja in sekcje:
        doc.add_heading(
            f"{TYTUL_SEKCJI_MODULU}: {sekcja['der_name']}",
            level=poziom_naglowka,
        )
        for etykieta, tresc, poziom in wiersze_sekcji(sekcja):
            akapit = doc.add_paragraph()
            akapit.paragraph_format.left_indent = Pt(14 * poziom)
            akapit.add_run(f"{etykieta}: ").bold = True
            tekst = akapit.add_run(tresc)
            tekst.bold = etykieta == POZYCJA_WYMAGANIE


__all__ = [
    "BRAK_CERTYFIKATU_PL",
    "POZYCJA_POMINIETY",
    "TYTUL_SEKCJI_MODULU",
    "dopisz_sekcje_docx",
    "opis_dowodu_certyfikatu",
    "opis_pominietego",
    "sekcja_modulu",
    "sekcje_modulow",
    "wiersze_dowodu",
    "wiersze_sekcji",
]
