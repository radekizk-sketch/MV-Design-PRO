"""
Reference Patterns API — Wzorce odniesienia

READ-ONLY API for running and viewing reference pattern validations.
Pattern A: Dobór I>> dla linii SN (selektywność, czułość, cieplne, SPZ)

CANONICAL ALIGNMENT:
- NOT-A-SOLVER: Reference patterns are INTERPRETATION layer (no physics)
- WHITE BOX: Full trace of validation steps
- DETERMINISM: Same inputs → identical outputs
- Polish labels (100% PL UI)

NO CODENAMES IN UI.
"""

from __future__ import annotations

import json
from typing import TYPE_CHECKING, Any
from uuid import uuid4

from api.odmowa_http import OdmowaHttpError
from application.reference_patterns import (
    PATTERN_ID,
    PATTERN_NAME_PL,
    ReferencePatternResult,
    fixture_to_input,
    get_pattern_a_fixtures_dir,
    run_pattern_a,
)
from fastapi import APIRouter, HTTPException, Response
from network_model.odmowa_danych import OdmowaDanychError
from network_model.reporting.czcionki import zarejestruj_czcionki
from pydantic import BaseModel
from werdykt import WyjasnienieWerdyktu, WynikWymagania

if TYPE_CHECKING:
    # python-docx jest importowany LENIWIE w ciele eksportu (brak paczki => 501),
    # wiec typ tabeli sprowadzamy wylacznie na potrzeby analizy statycznej.
    from docx.table import Table

router = APIRouter(
    prefix="/api/reference-patterns",
    tags=["reference-patterns"],
)


# =============================================================================
# Request/Response Models
# =============================================================================


class PatternMetadata(BaseModel):
    """Pattern metadata for listing."""

    pattern_id: str
    name_pl: str
    description_pl: str


class FixtureMetadata(BaseModel):
    """Fixture metadata for listing.

    ``expected_verdict`` — kod oczekiwany zapisany w fiksturze (przypadek odniesienia);
    ``wymaganie`` i ``wyjasnienie`` — rekord W wzorca uruchomionego na tej fiksturze (ten sam
    bieg co ``GET /fixtures/{plik}``), więc kod oczekiwany ma obok wyjaśnienie z liczbami.
    """

    fixture_id: str
    filename: str
    description: str | None = None
    expected_verdict: str | None = None
    notes_pl: str | None = None
    wymaganie: WynikWymagania
    wyjasnienie: WyjasnienieWerdyktu


class PatternRunRequest(BaseModel):
    """Request body for running a pattern."""

    pattern_id: str
    fixture_file: str | None = None
    input_override: dict[str, Any] | None = None


class PatternRunResponse(BaseModel):
    """Response for pattern run."""

    run_id: str
    pattern_id: str
    name_pl: str
    verdict: str
    verdict_description_pl: str
    summary_pl: str
    wymaganie: WynikWymagania
    wyjasnienie: WyjasnienieWerdyktu
    checks: list[dict[str, Any]]
    trace: list[dict[str, Any]]
    artifacts: dict[str, Any]


class PatternListResponse(BaseModel):
    """Response for pattern list."""

    patterns: list[PatternMetadata]


class FixtureListResponse(BaseModel):
    """Response for fixture list."""

    pattern_id: str
    fixtures: list[FixtureMetadata]


# =============================================================================
# Helper Functions
# =============================================================================


def result_to_response(result: ReferencePatternResult, run_id: str) -> PatternRunResponse:
    """Convert ReferencePatternResult to API response."""
    data = result.to_dict()
    return PatternRunResponse(
        run_id=run_id,
        pattern_id=data["pattern_id"],
        name_pl=data["name_pl"],
        verdict=data["verdict"],
        verdict_description_pl=data["verdict_description_pl"],
        summary_pl=data["summary_pl"],
        wymaganie=result.wymaganie,
        wyjasnienie=result.wyjasnienie,
        checks=data["checks"],
        trace=data["trace"],
        artifacts=data["artifacts"],
    )


def list_pattern_a_fixtures() -> list[FixtureMetadata]:
    """List available fixtures for Pattern A."""
    fixtures_dir = get_pattern_a_fixtures_dir()
    fixtures: list[FixtureMetadata] = []

    if not fixtures_dir.exists():
        return fixtures

    for filepath in sorted(fixtures_dir.glob("*.json")):
        # Fikstury wzorca są częścią produktu (karta #151): plik nieczytelny albo
        # niepoprawny JSON to defekt wydania, nie powód do cichego pominięcia pozycji
        # z listy — wybucha (500 z pełnym śladem w dzienniku).
        with open(filepath, encoding="utf-8") as f:
            data = json.load(f)

        fixture_id = filepath.stem
        wynik = run_pattern_a(input_data=fixture_to_input(data))
        fixtures.append(
            FixtureMetadata(
                fixture_id=fixture_id,
                filename=filepath.name,
                description=data.get("_description"),
                expected_verdict=data.get("_expected_verdict"),
                notes_pl=data.get("_notes_pl"),
                wymaganie=wynik.wymaganie,
                wyjasnienie=wynik.wyjasnienie,
            )
        )

    return fixtures


def _zawin(tekst: str, szerokosc: int) -> list[str]:
    """Zawinięcie tekstu raportu PDF do wierszy (bez obcinania treści wyjaśnienia)."""
    import textwrap

    return textwrap.wrap(tekst, width=szerokosc) or [""]


# =============================================================================
# Endpoints
# =============================================================================


@router.get("/patterns", response_model=PatternListResponse)
async def list_patterns() -> PatternListResponse:
    """
    List available reference patterns.

    Currently only Pattern A is available.
    """
    patterns = [
        PatternMetadata(
            pattern_id=PATTERN_ID,
            name_pl=PATTERN_NAME_PL,
            description_pl="Wzorzec walidacji doboru nastaw I>> dla linii SN — "
            "selektywność, czułość, wytrzymałość cieplna, blokada SPZ",
        )
    ]
    return PatternListResponse(patterns=patterns)


@router.get("/patterns/{pattern_id}/fixtures", response_model=FixtureListResponse)
def list_pattern_fixtures(pattern_id: str) -> FixtureListResponse:
    """
    List available fixtures (test cases) for a pattern.

    Pattern A fixtures: case_A_zgodne, case_B_niezgodne, case_C_graniczne
    """
    if pattern_id != PATTERN_ID:
        raise HTTPException(
            status_code=404,
            detail=f"Wzorzec '{pattern_id}' nie istnieje. Dostępny: {PATTERN_ID}",
        )

    fixtures = list_pattern_a_fixtures()
    return FixtureListResponse(pattern_id=pattern_id, fixtures=fixtures)


@router.post("/run", response_model=PatternRunResponse)
def run_pattern(request: PatternRunRequest) -> PatternRunResponse:
    """
    Run a reference pattern validation.

    Args:
        request: Pattern run request with pattern_id and optional fixture/input override

    Returns:
        Pattern validation result with verdict, checks, trace, and artifacts
    """
    if request.pattern_id != PATTERN_ID:
        raise HTTPException(
            status_code=400,
            detail=f"Nieznany wzorzec: '{request.pattern_id}'. Dostępny: {PATTERN_ID}",
        )

    run_id = str(uuid4())

    try:
        if request.fixture_file:
            # Run with fixture
            result = run_pattern_a(fixture_file=request.fixture_file)
        elif request.input_override:
            # Run with custom input
            # For now, we require a fixture base — override functionality can be added later
            raise HTTPException(
                status_code=400,
                detail="Uruchomienie z input_override wymaga bazowego fixture_file",
            )
        else:
            raise HTTPException(
                status_code=400,
                detail="Wymagany fixture_file lub input_override",
            )

        return result_to_response(result, run_id)

    except FileNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e)) from e
    except OdmowaDanychError as e:
        raise OdmowaHttpError(e, status_code=400) from e


@router.get("/fixtures/{fixture_file}", response_model=PatternRunResponse)
def run_pattern_with_fixture(fixture_file: str) -> PatternRunResponse:
    """
    Run Pattern A with a specific fixture file (convenience GET endpoint).

    Args:
        fixture_file: Fixture filename (e.g., "case_A_zgodne.json")

    Returns:
        Pattern validation result
    """
    run_id = str(uuid4())

    try:
        result = run_pattern_a(fixture_file=fixture_file)
        return result_to_response(result, run_id)
    except FileNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e)) from e
    except OdmowaDanychError as e:
        raise OdmowaHttpError(e, status_code=400) from e


# =============================================================================
# Export Endpoints — PDF and DOCX
# =============================================================================


@router.get("/fixtures/{fixture_file}/export/pdf")
def export_pattern_result_pdf(fixture_file: str) -> Response:
    """
    Export reference pattern result to PDF.

    Args:
        fixture_file: Fixture filename (e.g., "case_A_zgodne.json")

    Returns:
        PDF file as attachment
    """
    import io

    from fastapi.responses import Response

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise HTTPException(
            status_code=501,
            detail="Eksport PDF wymaga biblioteki reportlab. Zainstaluj: pip install reportlab",
        ) from exc

    try:
        result = run_pattern_a(fixture_file=fixture_file)
        data = result.to_dict()
    except FileNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e)) from e
    except OdmowaDanychError as e:
        raise OdmowaHttpError(e, status_code=400) from e

    # Create PDF in memory
    buffer = io.BytesIO()
    zarejestruj_czcionki()
    c = canvas.Canvas(buffer, pagesize=A4, invariant=1, pageCompression=0)
    page_width, page_height = A4

    left_margin = 25 * mm
    top_margin = page_height - 25 * mm
    y = top_margin
    line_height = 5 * mm

    # Title
    c.setFont("DejaVuSans-Bold", 16)
    title = f"Raport wzorca odniesienia: {data['name_pl']}"
    c.drawString(left_margin, y, title)
    y -= 10 * mm

    # Ocena wzorca — etykieta i zdanie rekordu W (kod werdyktu nie jest komunikatem)
    c.setFont("DejaVuSans-Bold", 14)
    c.drawString(left_margin, y, f"Ocena wzorca: {result.wymaganie.etykieta.etykieta_pl}")
    y -= 6 * mm

    c.setFont("DejaVuSans", 10)
    for wiersz in _zawin(data["verdict_description_pl"], 100) + _zawin(data["summary_pl"], 100):
        c.drawString(left_margin, y, wiersz)
        y -= line_height
        if y < 30 * mm:
            c.showPage()
            c.setFont("DejaVuSans", 10)
            y = top_margin
    y -= 5 * mm

    # Checks section
    c.setFont("DejaVuSans-Bold", 12)
    c.drawString(left_margin, y, "Sprawdzenia")
    y -= 6 * mm

    c.setFont("DejaVuSans", 9)
    for check in data.get("checks", []):
        tekst = f"{check['name_pl']} — {check['status_pl']}: {check['wyjasnienie']['zdanie_pl']}"
        for wiersz in _zawin(tekst, 110):
            c.drawString(left_margin, y, wiersz)
            y -= line_height
            if y < 30 * mm:
                c.showPage()
                c.setFont("DejaVuSans", 9)
                y = top_margin

    y -= 5 * mm

    # Artifacts section
    c.setFont("DejaVuSans-Bold", 12)
    c.drawString(left_margin, y, "Wartości pośrednie")
    y -= 6 * mm

    c.setFont("DejaVuSans", 9)
    artifacts = data.get("artifacts", {})
    artifact_items = [
        ("Czas zwarcia sumaryczny (tk)", artifacts.get("tk_total_s"), "s"),
        ("Prąd znamionowy cieplny (Ithn)", artifacts.get("ithn_a"), "A"),
        ("Prąd dopuszczalny cieplnie (Ithdop)", artifacts.get("ithdop_a"), "A"),
        ("I_min okna nastaw (pierwotna)", artifacts.get("window_i_min_primary_a"), "A"),
        ("I_max okna nastaw (pierwotna)", artifacts.get("window_i_max_primary_a"), "A"),
        ("Zalecana nastawa I>>", artifacts.get("recommended_setting_primary_a"), "A"),
        ("Okno nastaw prawidłowe", "Tak" if artifacts.get("window_valid") else "Nie", ""),
        ("Kryterium limitujące I_min", artifacts.get("limiting_criterion_min"), ""),
        ("Kryterium limitujące I_max", artifacts.get("limiting_criterion_max"), ""),
    ]

    for label, value, unit in artifact_items:
        if value is not None:
            if isinstance(value, float):
                text = f"{label}: {value:.4g} {unit}".strip()
            else:
                text = f"{label}: {value} {unit}".strip()
            c.drawString(left_margin, y, text)
            y -= line_height
            if y < 30 * mm:
                c.showPage()
                y = top_margin

    y -= 5 * mm

    # Trace section (summary)
    trace = data.get("trace", [])
    if trace:
        c.setFont("DejaVuSans-Bold", 12)
        c.drawString(left_margin, y, f"Ślad obliczeń ({len(trace)} kroków)")
        y -= 6 * mm

        c.setFont("DejaVuSans", 8)
        for step in trace[:20]:  # Limit to first 20 steps
            step_text = f"{step['step']}: {step['description_pl']}"
            if len(step_text) > 90:
                step_text = step_text[:87] + "..."
            c.drawString(left_margin, y, step_text)
            y -= line_height * 0.8
            if y < 30 * mm:
                c.showPage()
                y = top_margin

        if len(trace) > 20:
            c.drawString(left_margin, y, f"... oraz {len(trace) - 20} dodatkowych kroków")
            y -= line_height

    # Footer
    y = 15 * mm
    c.setFont("DejaVuSans", 8)
    c.drawString(left_margin, y, "Wygenerowano przez MV-DESIGN PRO — Wzorce odniesienia")

    c.save()
    buffer.seek(0)

    # Generate filename
    fixture_name = fixture_file.replace(".json", "")
    filename = f"wzorzec_odniesienia_{fixture_name}.pdf"

    return Response(
        content=buffer.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/fixtures/{fixture_file}/export/docx")
def export_pattern_result_docx(fixture_file: str) -> Response:
    """
    Export reference pattern result to DOCX.

    Args:
        fixture_file: Fixture filename (e.g., "case_A_zgodne.json")

    Returns:
        DOCX file as attachment

    Deterministyczny: ten sam wzorzec eksportowany wielokrotnie daje identyczne
    bajty (`docx_determinism.make_docx_bytes_deterministic` — sam python-docx
    NIE zeruje znacznikow czasu wpisow ZIP ani docProps/core.xml, dwa kolejne
    `Document().save()` roznia sie bajtowo, gdy wywolania przetna granice sekundy).
    """
    import io

    from fastapi.responses import Response
    from network_model.reporting.docx_determinism import make_docx_bytes_deterministic

    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError as exc:
        raise HTTPException(
            status_code=501,
            detail="Eksport DOCX wymaga biblioteki python-docx. Zainstaluj: pip install python-docx",
        ) from exc

    try:
        result = run_pattern_a(fixture_file=fixture_file)
        data = result.to_dict()
    except FileNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e)) from e
    except OdmowaDanychError as e:
        raise OdmowaHttpError(e, status_code=400) from e

    # Create DOCX document
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Title
    heading = doc.add_heading(f"Raport wzorca odniesienia: {data['name_pl']}", level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Ocena wzorca — etykieta i zdanie rekordu W (kod werdyktu nie jest komunikatem)
    verdict_para = doc.add_paragraph()
    verdict_run = verdict_para.add_run(f"Ocena wzorca: {result.wymaganie.etykieta.etykieta_pl}")
    verdict_run.bold = True
    verdict_run.font.size = Pt(14)

    doc.add_paragraph(data["verdict_description_pl"])
    doc.add_paragraph(data["summary_pl"])

    doc.add_paragraph()

    # Checks section
    doc.add_heading("Sprawdzenia", level=1)

    checks_table = doc.add_table(rows=1, cols=3)
    checks_table.style = "Table Grid"

    hdr = checks_table.rows[0].cells
    hdr[0].text = "Kryterium"
    hdr[1].text = "Ocena"
    hdr[2].text = "Wyjaśnienie"
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True

    for check in data.get("checks", []):
        row = checks_table.add_row().cells
        row[0].text = check["name_pl"]
        row[1].text = check["status_pl"]
        row[2].text = f"{check['wyjasnienie']['zdanie_pl']} {check['description_pl']}"

    doc.add_paragraph()

    # Artifacts section
    doc.add_heading("Wartości pośrednie", level=1)

    artifacts = data.get("artifacts", {})

    # Time and thermal
    doc.add_heading("Czas zwarcia i wytrzymałość cieplna", level=2)
    thermal_table = doc.add_table(rows=1, cols=2)
    thermal_table.style = "Table Grid"
    hdr = thermal_table.rows[0].cells
    hdr[0].text = "Parametr"
    hdr[1].text = "Wartosc"
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True

    def add_artifact_row(table: Table, label: str, value: Any, unit: str = "") -> None:
        if value is not None:
            row = table.add_row().cells
            row[0].text = label
            if isinstance(value, float):
                row[1].text = f"{value:.4g} {unit}".strip()
            else:
                row[1].text = f"{value} {unit}".strip()

    add_artifact_row(
        thermal_table, "Czas zwarcia sumaryczny (tk)", artifacts.get("tk_total_s"), "s"
    )
    add_artifact_row(thermal_table, "Prąd znamionowy cieplny (Ithn)", artifacts.get("ithn_a"), "A")
    add_artifact_row(
        thermal_table, "Prąd dopuszczalny cieplnie (Ithdop)", artifacts.get("ithdop_a"), "A"
    )

    doc.add_paragraph()

    # Window settings
    doc.add_heading("Okno nastaw I>>", level=2)
    window_table = doc.add_table(rows=1, cols=2)
    window_table.style = "Table Grid"
    hdr = window_table.rows[0].cells
    hdr[0].text = "Parametr"
    hdr[1].text = "Wartosc"
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True

    add_artifact_row(
        window_table,
        "I_min (selektywność, pierwotna)",
        artifacts.get("window_i_min_primary_a"),
        "A",
    )
    add_artifact_row(
        window_table, "I_max (okno, pierwotna)", artifacts.get("window_i_max_primary_a"), "A"
    )
    add_artifact_row(
        window_table,
        "Zalecana nastawa I>>",
        artifacts.get("recommended_setting_primary_a"),
        "A",
    )
    add_artifact_row(
        window_table, "Okno nastaw prawidłowe", "Tak" if artifacts.get("window_valid") else "Nie"
    )
    add_artifact_row(
        window_table, "Kryterium limitujące I_min", artifacts.get("limiting_criterion_min")
    )
    add_artifact_row(
        window_table, "Kryterium limitujące I_max", artifacts.get("limiting_criterion_max")
    )

    doc.add_paragraph()

    # Trace section
    trace = data.get("trace", [])
    if trace:
        doc.add_heading(f"Ślad obliczeń ({len(trace)} kroków)", level=1)

        for step in trace[:30]:  # Limit to first 30 steps
            step_para = doc.add_paragraph()
            step_run = step_para.add_run(f"{step['step']}: ")
            step_run.bold = True
            step_para.add_run(step["description_pl"])

        if len(trace) > 30:
            doc.add_paragraph(f"... oraz {len(trace) - 30} dodatkowych kroków")

    doc.add_paragraph()

    # Footer
    footer_para = doc.add_paragraph()
    footer_run = footer_para.add_run("Wygenerowano przez MV-DESIGN PRO — Wzorce odniesienia")
    footer_run.font.size = Pt(9)
    footer_run.italic = True

    # Save to BytesIO
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    # Generate filename
    fixture_name = fixture_file.replace(".json", "")
    filename = f"wzorzec_odniesienia_{fixture_name}.docx"

    return Response(
        content=make_docx_bytes_deterministic(buffer.getvalue()),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
