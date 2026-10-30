"""API tests for `api/reference_patterns.py` — wzorzec RP-LINE-I2-THERMAL-SPZ.

Karta W3-C2 (2026-09-09): wzorzec przebudowany na `ProtectionSettingsEngine`
(Hoppel/IRiESD). Ten plik testuje END-TO-END przez prawdziwy HTTP endpoint
(nie tylko warstwę Pythona) — POST /run, GET /fixtures/{plik}, listing
wzorcow/fixture'ow, eksport PDF/DOCX (status 200, niepusta tresc) — dla
wszystkich czterech fixture'ow wzorca (ZGODNE/NIEZGODNE/GRANICZNE/generacja
lokalna), zeby rozszerzenie generacji lokalnej i okna nastaw byly widoczne
takze przez API, nie tylko w warstwie Pythona.
"""

from __future__ import annotations

from fastapi.testclient import TestClient

FIXTURE_FILES = (
    "case_A_zgodne.json",
    "case_B_niezgodne_konflikt.json",
    "case_C_graniczne_waskie_okno.json",
    "case_D_generacja_lokalna.json",
)


def _client() -> TestClient:
    from api.main import app

    return TestClient(app)


def test_list_patterns_returns_pattern_a() -> None:
    response = _client().get("/api/reference-patterns/patterns")

    assert response.status_code == 200, response.text
    data = response.json()
    pattern_ids = [p["pattern_id"] for p in data["patterns"]]
    assert "RP-LINE-I2-THERMAL-SPZ" in pattern_ids


def test_list_fixtures_returns_all_four_cases() -> None:
    response = _client().get("/api/reference-patterns/patterns/RP-LINE-I2-THERMAL-SPZ/fixtures")

    assert response.status_code == 200, response.text
    data = response.json()
    filenames = {f["filename"] for f in data["fixtures"]}
    assert filenames == set(FIXTURE_FILES)


def test_get_fixture_run_returns_expected_verdict_case_a() -> None:
    response = _client().get("/api/reference-patterns/fixtures/case_A_zgodne.json")

    assert response.status_code == 200, response.text
    data = response.json()
    assert data["pattern_id"] == "RP-LINE-I2-THERMAL-SPZ"
    assert data["verdict"] == "ZGODNE"
    assert data["artifacts"]["window_valid"] is True
    # Klucz renamed od karty W3-C2 (Hoppel nie ma strony wtornej/przekladni CT)
    assert "recommended_setting_primary_a" in data["artifacts"]
    assert "recommended_setting_secondary_a" not in data["artifacts"]


def test_get_fixture_run_returns_expected_verdict_case_b_niezgodne() -> None:
    response = _client().get("/api/reference-patterns/fixtures/case_B_niezgodne_konflikt.json")

    assert response.status_code == 200, response.text
    data = response.json()
    assert data["verdict"] == "NIEZGODNE"
    assert data["artifacts"]["window_valid"] is False
    assert data["artifacts"]["window_conflict_pl"] is not None


def test_get_fixture_run_returns_expected_verdict_case_c_graniczne() -> None:
    response = _client().get("/api/reference-patterns/fixtures/case_C_graniczne_waskie_okno.json")

    assert response.status_code == 200, response.text
    data = response.json()
    assert data["verdict"] == "GRANICZNE"


def test_get_fixture_run_returns_expected_verdict_case_d_generacja_lokalna() -> None:
    """Rozszerzenie W3-C2 widoczne przez API: ryzyko blokady ZSZ od E-L."""
    response = _client().get("/api/reference-patterns/fixtures/case_D_generacja_lokalna.json")

    assert response.status_code == 200, response.text
    data = response.json()
    assert data["verdict"] == "GRANICZNE"
    assert data["artifacts"]["generacja_lokalna_aktywna"] is True
    assert data["artifacts"]["generacja_lokalna_ryzyko_zsz"] is True
    lg_check = next(c for c in data["checks"] if c["name_pl"] == "Generacja lokalna (E-L)")
    assert lg_check["status"] == "WARN"


def test_post_run_with_fixture_file_matches_get_endpoint() -> None:
    client = _client()
    via_get = client.get("/api/reference-patterns/fixtures/case_A_zgodne.json")
    via_post = client.post(
        "/api/reference-patterns/run",
        json={"pattern_id": "RP-LINE-I2-THERMAL-SPZ", "fixture_file": "case_A_zgodne.json"},
    )

    assert via_get.status_code == 200, via_get.text
    assert via_post.status_code == 200, via_post.text
    get_data = via_get.json()
    post_data = via_post.json()
    for klucz in ("pattern_id", "name_pl", "verdict", "summary_pl", "checks", "artifacts"):
        assert get_data[klucz] == post_data[klucz]


def test_post_run_unknown_pattern_id_returns_400() -> None:
    response = _client().post(
        "/api/reference-patterns/run",
        json={"pattern_id": "RP-NIEISTNIEJE", "fixture_file": "case_A_zgodne.json"},
    )
    assert response.status_code == 400


def test_get_fixture_run_unknown_fixture_returns_404() -> None:
    response = _client().get("/api/reference-patterns/fixtures/brak_takiego_pliku.json")
    assert response.status_code == 404


def test_export_pdf_returns_200_nonempty_for_all_fixtures() -> None:
    client = _client()
    for fixture_file in FIXTURE_FILES:
        response = client.get(f"/api/reference-patterns/fixtures/{fixture_file}/export/pdf")
        assert response.status_code == 200, f"{fixture_file}: {response.text}"
        assert response.content[:4] == b"%PDF", f"{fixture_file}: not a PDF"
        assert len(response.content) > 500, f"{fixture_file}: suspiciously small PDF"


def test_export_docx_returns_200_nonempty_for_all_fixtures() -> None:
    client = _client()
    for fixture_file in FIXTURE_FILES:
        response = client.get(f"/api/reference-patterns/fixtures/{fixture_file}/export/docx")
        assert response.status_code == 200, f"{fixture_file}: {response.text}"
        assert response.content[:2] == b"PK", f"{fixture_file}: not a ZIP/DOCX"
        assert len(response.content) > 500, f"{fixture_file}: suspiciously small DOCX"


# --- Werdykt wyjaśnialny (karta WERDYKT-MIGRACJA-190, grupa 5) -------------------------------
# Iloczyn cech: {lista fikstur, bieg fikstury, eksport PDF, eksport DOCX} × {cztery fikstury}.
# Kod `verdict` jest projekcją rekordu W; każdy kanał niesie rekord albo jego etykietę.


def test_lista_fikstur_niesie_rekord_zgodny_z_biegiem() -> None:
    client = _client()
    lista = client.get("/api/reference-patterns/patterns/RP-LINE-I2-THERMAL-SPZ/fixtures")
    assert lista.status_code == 200, lista.text
    for meta in lista.json()["fixtures"]:
        bieg = client.get(f"/api/reference-patterns/fixtures/{meta['filename']}").json()
        assert meta["wymaganie"] == bieg["wymaganie"], meta["filename"]
        assert meta["wyjasnienie"] == bieg["wymaganie"]["wyjasnienie"]
        # Przypadek odniesienia: kod oczekiwany zapisany w fiksturze = kod policzony z rekordu.
        assert meta["expected_verdict"] == bieg["verdict"], meta["filename"]


def test_bieg_fikstury_sprawdzenia_z_etykieta_rekordu() -> None:
    client = _client()
    for fixture_file in FIXTURE_FILES:
        data = client.get(f"/api/reference-patterns/fixtures/{fixture_file}").json()
        wymaganie = data["wymaganie"]
        assert data["wyjasnienie"] == wymaganie["wyjasnienie"]
        assert data["verdict_description_pl"].startswith(wymaganie["wyjasnienie"]["zdanie_pl"])
        kryteria_w = {o["kryterium_id"] for o in wymaganie["oceny_skladowe"]}
        for check in data["checks"]:
            assert check["status_pl"] == check["ocena"]["etykieta"]["etykieta_pl"]
            assert check["wyjasnienie"] == check["ocena"]["wyjasnienie"]
        # Diagnostyka generacji lokalnej stoi obok rekordu W, nie w nim.
        assert len(kryteria_w) == len(data["checks"]) - 1


def test_eksport_pdf_i_docx_niesie_etykiete_rekordu() -> None:
    import io

    import pymupdf
    from docx import Document

    client = _client()
    for fixture_file in FIXTURE_FILES:
        data = client.get(f"/api/reference-patterns/fixtures/{fixture_file}").json()
        etykieta = data["wymaganie"]["etykieta"]["etykieta_pl"]

        pdf = client.get(f"/api/reference-patterns/fixtures/{fixture_file}/export/pdf")
        assert pdf.status_code == 200
        with pymupdf.open(stream=pdf.content, filetype="pdf") as dokument:
            tekst_pdf = " ".join(" ".join(strona.get_text().split()) for strona in dokument)
        assert f"Ocena wzorca: {etykieta}" in tekst_pdf, fixture_file
        for check in data["checks"]:
            assert f"{check['name_pl']} — {check['status_pl']}" in tekst_pdf, check["name_pl"]

        docx = client.get(f"/api/reference-patterns/fixtures/{fixture_file}/export/docx")
        assert docx.status_code == 200
        dokument_docx = Document(io.BytesIO(docx.content))
        tekst_docx = "\n".join(p.text for p in dokument_docx.paragraphs)
        komorki = {k.text for t in dokument_docx.tables for w in t.rows for k in w.cells}
        assert "Ocena wzorca:" in tekst_docx or "Ocena wzorca:" in komorki, fixture_file
        assert etykieta in tekst_docx or etykieta in komorki, fixture_file
        for check in data["checks"]:
            assert check["status_pl"] in komorki, check["name_pl"]
