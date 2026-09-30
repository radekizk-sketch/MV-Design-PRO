"""Kontrakt API koordynacji zabezpieczeń E-28 (karta BIEG-ZABEZPIECZEN-Z-MODELU).

Żądanie niesie wyłącznie identyfikatory biegów (MAX, MIN, rozpływ), opcjonalne pary i
kryteria. Urządzenia i nastawy — z BIEŻĄCEGO modelu projektu (sieć złota G08 w magazynie pod
kluczem projektu), prądy — z biegów. Iloczyn cech:
  trasa {run, wynik, pdf, docx} × {wynik istnieje, nieznany run_id},
  odmowa {bieg niewskazany, bieg nieistniejący, scenariusze zamienione, model bez
  zabezpieczeń, sieć zmieniona od biegu, pola dawnego kontraktu (devices/prądy)},
  eksport {PDF, DOCX} × determinizm bajtowy × nazwy urządzeń z modelu.
"""

from __future__ import annotations

import hashlib
from collections.abc import Iterator
from typing import Any
from uuid import UUID

import pytest

pytest.importorskip("fastapi")

from enm.canonical_analysis import create_run, execute_run, reset_canonical_runs  # noqa: E402
from enm.domain_operations import execute_domain_operation  # noqa: E402
from enm.klucz_twin import klucz_twin_projektu  # noqa: E402
from enm.models import EnergyNetworkModel  # noqa: E402
from enm.store import get_enm, set_enm  # noqa: E402

from tests.golden.enm_builders.zabezpieczenia_magistrali import (  # noqa: E402
    NAZWA_ZABEZPIECZENIA_Q1,
    NAZWA_ZABEZPIECZENIA_Q2,
    build_zabezpieczenia_magistrali_enm,
    siec_magistrali,
)


def _projekt(app_client: Any, model: EnergyNetworkModel) -> tuple[str, dict[str, str]]:
    odpowiedz = app_client.post("/api/projects", json={"name": "Projekt koordynacji"})
    assert odpowiedz.status_code == 201, odpowiedz.text
    projekt = odpowiedz.json()["id"]
    klucz = klucz_twin_projektu(UUID(projekt))
    set_enm(klucz, model)
    biegi: dict[str, str] = {}
    for nazwa, rodzaj, opcje in (
        ("max", "short_circuit_sn", None),
        ("min", "short_circuit_sn", {"scenario": "min"}),
        ("pf", "PF", None),
    ):
        bieg = execute_run(
            create_run(
                case_id=f"case-{nazwa}",
                klucz_twin=klucz,
                analysis_type=rodzaj,
                project_id=projekt,
                options=opcje,
            ).id
        )
        assert bieg.status == "FINISHED", bieg.error_message
        biegi[nazwa] = str(bieg.id)
    return projekt, biegi


@pytest.fixture()
def scena(app_client: Any) -> Iterator[tuple[Any, str, dict[str, str]]]:
    reset_canonical_runs()
    projekt, biegi = _projekt(app_client, build_zabezpieczenia_magistrali_enm())
    yield app_client, projekt, biegi
    reset_canonical_runs()


def _uruchom(app_client: Any, projekt: str, cialo: dict[str, Any]) -> Any:
    return app_client.post(f"/api/protection-coordination/projects/{projekt}/run", json=cialo)


def _zadanie(biegi: dict[str, str]) -> dict[str, Any]:
    return {"sc_run_id": biegi["max"], "sc_run_id_min": biegi["min"], "pf_run_id": biegi["pf"]}


def test_koordynacja_urzadzen_modelu_przez_api(scena: tuple[Any, str, dict[str, str]]) -> None:
    klient, projekt, biegi = scena
    odpowiedz = _uruchom(klient, projekt, _zadanie(biegi))
    assert odpowiedz.status_code == 201, odpowiedz.text
    podsumowanie = odpowiedz.json()
    # P-06: potwierdzenie niesie liczby zbiorcze, nie werdykt ogólny.
    assert "overall_verdict" not in podsumowanie
    assert podsumowanie["total_devices"] == 2
    assert podsumowanie["total_checks"] == 5
    assert podsumowanie["najmniejszy_odstep_s"] == pytest.approx(0.3)
    assert podsumowanie["najmniejszy_iloraz_czulosci"] > 1.0
    assert podsumowanie["najmniejszy_iloraz_przeciazalnosci"] > 1.0

    run_id = podsumowanie["run_id"]
    wynik = klient.get(f"/api/protection-coordination/{run_id}").json()
    assert {d["name"] for d in wynik["devices"]} == {
        NAZWA_ZABEZPIECZENIA_Q1,
        NAZWA_ZABEZPIECZENIA_Q2,
    }
    assert all(d["nastawy"]["stopnie"] for d in wynik["devices"])
    assert len(wynik["pary"]) == 1
    assert len(wynik["tcc_curves"]) == 2 and all(k["points"] for k in wynik["tcc_curves"])
    assert wynik["fault_markers"]
    assert wynik["trace_steps"]
    for rodzaj, liczba in (("sensitivity", 2), ("selectivity", 1), ("overload", 2)):
        assert len(wynik[f"{rodzaj}_checks"]) == liczba
        assert all("verdict" not in c for c in wynik[f"{rodzaj}_checks"])
    assert "overall_verdict" not in wynik


def test_kryteria_spoza_kontraktu_odrzucone(scena: tuple[Any, str, dict[str, str]]) -> None:
    """Dawne progi pasm werdyktu (``*_marginal``, ``cti_margin_factor``) nie są kryterium —
    żądanie z nimi jest odrzucane jawnie, nie ignorowane po cichu."""
    klient, projekt, biegi = scena
    odpowiedz = _uruchom(
        klient,
        projekt,
        {**_zadanie(biegi), "config": {"sensitivity_margin_marginal": 1.2}},
    )
    assert odpowiedz.status_code == 422


def test_podtrasy_wycinkow_wyniku_skasowane(app_client: Any) -> None:
    """Podtrasy `/tcc`, `/trace`, `/checks/*` bez konsumenta — skasowane (jeden odczyt)."""
    sciezki = {r.path for r in app_client.app.routes}
    for sufiks in (
        "/tcc",
        "/trace",
        "/checks/sensitivity",
        "/checks/selectivity",
        "/checks/overload",
    ):
        assert f"/api/protection-coordination/{{run_id}}{sufiks}" not in sciezki


@pytest.mark.parametrize(
    ("pola", "powod"),
    [
        ({"sc_run_id": None}, "BIEG_NIE_WSKAZANY"),
        ({"sc_run_id": "00000000-0000-0000-0000-000000000000"}, "BIEG_NIE_ISTNIEJE"),
    ],
)
def test_odmowy_mostu_autorytetu(
    scena: tuple[Any, str, dict[str, str]], pola: dict[str, Any], powod: str
) -> None:
    klient, projekt, biegi = scena
    odpowiedz = _uruchom(klient, projekt, {**_zadanie(biegi), **pola})
    assert odpowiedz.status_code == 422
    assert odpowiedz.json()["detail"]["powod"] == powod


def test_scenariusze_zamienione_odmowa(scena: tuple[Any, str, dict[str, str]]) -> None:
    klient, projekt, biegi = scena
    odpowiedz = _uruchom(
        klient,
        projekt,
        {"sc_run_id": biegi["min"], "sc_run_id_min": biegi["max"], "pf_run_id": biegi["pf"]},
    )
    assert odpowiedz.status_code == 422
    assert odpowiedz.json()["detail"]["powod"] == "SCENARIUSZ_BIEGU_NIEZGODNY"


def test_siec_zmieniona_od_biegu_odmowa(scena: tuple[Any, str, dict[str, str]]) -> None:
    klient, projekt, biegi = scena
    klucz = klucz_twin_projektu(UUID(projekt))
    dane = get_enm(klucz).model_dump(mode="json")
    kabel = next(g["ref_id"] for g in dane["branches"] if g["type"] == "cable")
    wynik = execute_domain_operation(
        dane,
        "update_element_parameters",
        {"element_ref": kabel, "parameters": {"length_km": 4.0, "parameter_source": "CATALOG"}},
    )
    assert wynik.get("error") is None
    set_enm(klucz, EnergyNetworkModel.model_validate(wynik["snapshot"]))
    odpowiedz = _uruchom(klient, projekt, _zadanie(biegi))
    assert odpowiedz.status_code == 422
    assert odpowiedz.json()["detail"]["powod"] == "SIEC_ZMIENIONA_OD_BIEGU"


def test_model_bez_zabezpieczen_odmowa(app_client: Any) -> None:
    reset_canonical_runs()
    enm, _a, _b = siec_magistrali()
    projekt, biegi = _projekt(app_client, EnergyNetworkModel.model_validate(enm))
    odpowiedz = _uruchom(app_client, projekt, _zadanie(biegi))
    assert odpowiedz.status_code == 422
    assert odpowiedz.json()["detail"]["powod"] == "BRAK_ZABEZPIECZEN_W_MODELU"


@pytest.mark.parametrize(
    "pole",
    ["devices", "fault_currents", "operating_currents"],
)
def test_pola_dawnego_kontraktu_odrzucone(
    scena: tuple[Any, str, dict[str, str]], pole: str
) -> None:
    """Urządzenia, prądy zwarciowe i robocze od klienta skasowane — ładunek z nimi jest
    błędem walidacji (``extra=forbid``), nie cichym zignorowaniem."""
    klient, projekt, biegi = scena
    odpowiedz = _uruchom(klient, projekt, {**_zadanie(biegi), pole: []})
    assert odpowiedz.status_code == 422


@pytest.mark.parametrize(
    "sufiks",
    ["", "/export/pdf", "/export/docx"],
)
def test_nieznany_run_id_404(app_client: Any, sufiks: str) -> None:
    odpowiedz = app_client.get(f"/api/protection-coordination/nieznany-bieg{sufiks}")
    assert odpowiedz.status_code == 404


@pytest.mark.parametrize(
    ("sufiks", "typ"),
    [
        ("pdf", "application/pdf"),
        ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ],
)
def test_eksport_deterministyczny(
    scena: tuple[Any, str, dict[str, str]], sufiks: str, typ: str
) -> None:
    pytest.importorskip("reportlab" if sufiks == "pdf" else "docx")
    klient, projekt, biegi = scena
    run_id = _uruchom(klient, projekt, _zadanie(biegi)).json()["run_id"]
    pierwszy = klient.get(f"/api/protection-coordination/{run_id}/export/{sufiks}")
    drugi = klient.get(f"/api/protection-coordination/{run_id}/export/{sufiks}")
    assert pierwszy.status_code == 200
    assert pierwszy.headers["content-type"].startswith(typ)
    assert "attachment" in pierwszy.headers["content-disposition"]
    assert hashlib.sha256(pierwszy.content).digest() == hashlib.sha256(drugi.content).digest()


def test_eksport_docx_nazywa_urzadzenia_z_modelu(scena: tuple[Any, str, dict[str, str]]) -> None:
    pytest.importorskip("docx")
    import io

    from docx import Document

    klient, projekt, biegi = scena
    run_id = _uruchom(klient, projekt, _zadanie(biegi)).json()["run_id"]
    tresc = klient.get(f"/api/protection-coordination/{run_id}/export/docx").content
    dokument = Document(io.BytesIO(tresc))
    komorki = {c.text for t in dokument.tables for w in t.rows for c in w.cells}
    assert NAZWA_ZABEZPIECZENIA_Q1 in komorki
    assert NAZWA_ZABEZPIECZENIA_Q2 in komorki
    assert "Brak urządzeń" not in "\n".join(p.text for p in dokument.paragraphs)
