"""Spójność zbioru gałęzi rozpływu zwarciowego na REALNYCH ścieżkach (karta PROG-ZERA-WKLADOW,
wymaganie 4 właściciela).

Ścieżka użytkownika: projekt + przypadek przez API → bieg zwarciowy (`POST …/runs` +
wykonanie) → wiersze wyniku → rozpływ punktu (`GET …/short-circuit/rozplyw`) → eksporty
raportu (JSON, DOCX, PDF) i pakietu dowodowego (JSON, LaTeX, PDF).

Iloczyn cech: rodzaj zwarcia {3F, 2F, 1F, 2F+Z} × punkt zwarcia (każdy węzeł sieci) ×
miejsce prezentacji {lista wkładów API, ślad WHITE BOX, zestawienie wkładów pominiętych,
eksport raportu, pakiet dowodowy}. Ten sam zbiór gałęzi znaczy:
- lista wkładów i zestawienie wkładów pominiętych są ROZŁĄCZNE i razem dają każdą parę
  źródło×gałąź policzoną przez rdzeń;
- kroki śladu podziału prądu Thevenina nazywają DOKŁADNIE gałęzie wkładów sieci zastępczej
  z listy (wkład pominięty nie ma kroku), a krok `wklady_pominiete` liczy DOKŁADNIE wpisy
  zestawienia per tor;
- eksporty raportu i pakiet dowodowy nie niosą rozpływu gałęziowego w ŻADNEJ postaci
  (decyzja K13: rozpływ jednego punktu wyłącznie endpointem na żądanie; wiersz niesie
  flagę dostępności) — więc nie mogą pokazać innego zbioru niż API.
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from typing import Any

import pytest
from api.main import app
from fastapi.testclient import TestClient
from network_model.solvers.short_circuit_contributions import PROG_WKLADU_POMIJALNEGO_A

from tests.utils.biegi import wykonaj_bieg
from tests.utils.siec_wkladow_pominietych import (
    NAZWA_GALEZI_BEZ_PRADU,
    NAZWA_GALEZI_REALNEJ_MALEJ,
    siec_wkladow_pominietych,
)

KLUCZE_ROZPLYWU_W_EKSPORCIE = ("branch_flow_trace", "wklady_pominiete")


@pytest.fixture
def client() -> Iterator[TestClient]:
    from api.power_flow_runs import _interpretation_cache
    from enm.canonical_analysis import reset_canonical_runs
    from enm.store import reset_enm_store

    _interpretation_cache.clear()
    reset_canonical_runs()
    reset_enm_store()
    with TestClient(app) as test_client:
        yield test_client


def _przypadek_z_siecia(client: TestClient) -> str:
    from application.twin_key import klucz_twin_dla_przypadku
    from enm.store import set_enm

    projekt = client.post("/api/projects", json={"name": "Prog wkladu pominietego"})
    assert projekt.status_code == 201, projekt.text
    przypadek = client.post(
        "/api/study-cases", json={"project_id": projekt.json()["id"], "name": "Przypadek"}
    )
    assert przypadek.status_code == 201, przypadek.text
    case_id = str(przypadek.json()["id"])
    set_enm(
        klucz_twin_dla_przypadku(case_id, client.app.state.uow_factory),
        siec_wkladow_pominietych(),
    )
    return case_id


def _liscie(obiekt: Any, klucz: str | None = None) -> Iterator[tuple[str | None, Any]]:
    if isinstance(obiekt, dict):
        for k, v in obiekt.items():
            yield k, v
            yield from _liscie(v, k)
    elif isinstance(obiekt, list):
        for v in obiekt:
            yield from _liscie(v, klucz)


def _tekst_docx(tresc: bytes) -> str:
    import docx

    dokument = docx.Document(io.BytesIO(tresc))
    akapity = [p.text for p in dokument.paragraphs]
    komorki = [c.text for t in dokument.tables for r in t.rows for c in r.cells]
    return "\n".join(akapity + komorki)


def _tekst_pdf(tresc: bytes) -> str:
    import pymupdf

    with pymupdf.open(stream=tresc, filetype="pdf") as dokument:
        return "\n".join(strona.get_text() for strona in dokument)


@pytest.mark.parametrize("rodzaj", ["SC_3F", "SC_2F", "SC_1F", "SC_2F_G"])
def test_ten_sam_zbior_galezi_w_api_sladzie_i_eksportach(client: TestClient, rodzaj: str) -> None:
    case_id = _przypadek_z_siecia(client)
    utworz = client.post(
        f"/api/execution/study-cases/{case_id}/runs",
        json={"analysis_type": rodzaj, "solver_input": {}},
    )
    assert utworz.status_code == 201, utworz.text
    run_id = utworz.json()["id"]
    wykonanie = wykonaj_bieg(client, run_id)
    assert wykonanie.status_code == 200 and wykonanie.json()["status"] == "DONE"

    wiersze = client.get(f"/api/analysis-runs/{run_id}/results/short-circuit").json()["rows"]
    assert len(wiersze) == 4
    widziane_pominiete: set[str] = set()
    widziane_realne_male: set[str] = set()
    for wiersz in wiersze:
        assert wiersz["branch_contributions"] is None
        assert wiersz["branch_contributions_available"] is True
        odpowiedz = client.get(
            f"/api/analysis-runs/{run_id}/results/short-circuit/rozplyw",
            params={"target_id": wiersz["target_id"]},
        )
        assert odpowiedz.status_code == 200
        rozplyw = odpowiedz.json()
        # Próg jako ODBICIE jedynej definicji (frontend nie ma własnej stałej).
        assert rozplyw["prog_wkladu_pominietego_a"] == PROG_WKLADU_POMIJALNEGO_A
        wklady = rozplyw["branch_contributions"]
        pominiete = rozplyw["wklady_pominiete"] or []
        slad = rozplyw["branch_flow_trace"]

        pary_wkladow = {(w["source_id"], w["branch_id"]) for w in wklady}
        pary_pominietych = {(w["source_id"], w["branch_id"]) for w in pominiete}
        assert len(pary_wkladow) == len(wklady)
        assert len(pary_pominietych) == len(pominiete)
        assert not pary_wkladow & pary_pominietych
        for w in wklady:
            assert w["i_ka"] * 1000.0 >= PROG_WKLADU_POMIJALNEGO_A
            assert w["direction"] in {"from_to", "to_from"}
        for w in pominiete:
            assert w["i_ka"] * 1000.0 < PROG_WKLADU_POMIJALNEGO_A
            assert w["tor"] in {"thevenin", "falownik"}
        # Każda para źródło×gałąź policzona przez rdzeń jest w DOKŁADNIE jednym zbiorze:
        # sieć ma 3 gałęzie; Thevenin zawsze, falownik w 3F/2F (PV bez składowej zerowej).
        zrodla = {s for s, _b in pary_wkladow | pary_pominietych}
        galezie = {b for _s, b in pary_wkladow | pary_pominietych}
        assert len(galezie) == 3
        assert pary_wkladow | pary_pominietych == {(s, b) for s in zrodla for b in galezie}

        # Ślad ↔ lista: kroki Thevenina nazywają DOKŁADNIE gałęzie wkładów sieci zastępczej.
        nazwy_thevenin = {w["branch_name"] for w in wklady if w["source_id"] == "THEVENIN_GRID"}
        nazwy_krokow = {
            k["title"].removeprefix("Prąd zwarciowy Thevenina w gałęzi ")
            for k in slad
            if k["key"].startswith("thevenin_flow_")
            and k["key"] not in ("thevenin_flow_setup", "thevenin_flow_balance")
        }
        assert nazwy_krokow == nazwy_thevenin
        kroki_pominietych = [k for k in slad if k["key"] == "wklady_pominiete"]
        assert len(kroki_pominietych) == (1 if pominiete else 0)
        if pominiete:
            wynik = kroki_pominietych[0]["result"]
            for tor in ("thevenin", "falownik"):
                liczba = sum(1 for w in pominiete if w["tor"] == tor)
                assert wynik.get(f"liczba_pominietych_{tor}", 0) == liczba
        widziane_pominiete |= {w["branch_name"] for w in pominiete}
        widziane_realne_male |= {
            w["branch_name"]
            for w in wklady
            if w["source_id"] == "THEVENIN_GRID"
            and w["i_ka"] * 1000.0 < 1e-4
            and w["branch_name"] == NAZWA_GALEZI_REALNEJ_MALEJ
        }
    # Wpis pominięty istnieje (gałąź martwa bez pojemności), a realny mały prąd ładowania
    # odgałęzienia (kilka µA, nad progiem) ZOSTAJE na liście.
    assert NAZWA_GALEZI_BEZ_PRADU in widziane_pominiete
    assert widziane_realne_male == {NAZWA_GALEZI_REALNEJ_MALEJ}

    # Eksporty: ŻADNEJ postaci rozpływu gałęziowego (K13) — zbiór pusty, zgodny z API.
    raport = client.get(f"/api/analysis-runs/{run_id}/export/report/json")
    assert raport.status_code == 200
    dowod = client.get(f"/api/analysis-runs/{run_id}/export/proof/json")
    assert dowod.status_code == 200
    for tresc in (raport.json(), dowod.json()):
        for klucz, wartosc in _liscie(tresc):
            if klucz == "branch_contributions":
                assert wartosc is None
            assert klucz not in KLUCZE_ROZPLYWU_W_EKSPORCIE
            if isinstance(wartosc, str):
                assert not wartosc.startswith("thevenin_flow_")
    for wiersz in raport.json()["results"]["short_circuit"]["rows"]:
        assert wiersz["branch_contributions_available"] is True
    teksty = {
        "raport DOCX": _tekst_docx(
            client.get(f"/api/analysis-runs/{run_id}/export/report/docx").content
        ),
        "raport PDF": _tekst_pdf(
            client.get(f"/api/analysis-runs/{run_id}/export/report/pdf").content
        ),
        "dowód PDF": _tekst_pdf(
            client.get(f"/api/analysis-runs/{run_id}/export/proof/pdf").content
        ),
        "dowód LaTeX": client.get(f"/api/analysis-runs/{run_id}/export/proof/latex").text,
    }
    for nazwa, tekst in teksty.items():
        assert tekst, nazwa
        assert "Prąd zwarciowy Thevenina w gałęzi" not in tekst, nazwa
        assert "pominięt" not in tekst, nazwa
